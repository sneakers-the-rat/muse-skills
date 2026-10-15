#!/usr/bin/env python3
"""Deterministic local DOCX extraction pipeline for the documents skill.

Outputs:
- manifest.json
- segments.jsonl
- chunks.jsonl
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import sys
from pathlib import Path
from typing import Iterable


def normalize_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def iter_segments(document) -> Iterable[dict]:
    segment_id = 1
    for paragraph_index, paragraph in enumerate(document.paragraphs, start=1):
        text = normalize_text(paragraph.text or "")
        if not text:
            continue
        yield {
            "segment_id": segment_id,
            "kind": "paragraph",
            "paragraph_index": paragraph_index,
            "char_count": len(text),
            "text": text,
        }
        segment_id += 1

    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            cells = [normalize_text(cell.text or "") for cell in row.cells]
            cells = [cell for cell in cells if cell]
            if not cells:
                continue
            text = " | ".join(cells)
            yield {
                "segment_id": segment_id,
                "kind": "table_row",
                "table_index": table_index,
                "row_index": row_index,
                "column_count": len(cells),
                "char_count": len(text),
                "text": text,
            }
            segment_id += 1


def iter_chunks(
    segments: list[dict], chunk_chars: int, overlap_chars: int
) -> Iterable[dict]:
    chunk_id = 1
    i = 0
    while i < len(segments):
        start_segment = segments[i]["segment_id"]
        buf: list[str] = []
        end_segment = start_segment
        char_count = 0
        j = i
        while j < len(segments):
            segment_text = segments[j]["text"]
            extra = len(segment_text) + (2 if buf else 0)
            if buf and char_count + extra > chunk_chars:
                break
            if buf:
                buf.append("\n\n")
            buf.append(segment_text)
            char_count += extra
            end_segment = segments[j]["segment_id"]
            j += 1

        text = "".join(buf).strip()
        if text:
            yield {
                "chunk_id": chunk_id,
                "start_segment": start_segment,
                "end_segment": end_segment,
                "char_count": len(text),
                "text": text,
            }
            chunk_id += 1

        if j >= len(segments):
            break

        if overlap_chars <= 0:
            i = j
            continue

        rewind = 0
        k = j - 1
        while k >= i and rewind < overlap_chars:
            rewind += len(segments[k]["text"])
            k -= 1
        i = max(i + 1, k + 1)


def main() -> int:
    parser = argparse.ArgumentParser(description="Extract deterministic text from DOCX files.")
    parser.add_argument("--input", required=True, help="DOCX path")
    parser.add_argument("--outdir", required=True, help="Output directory")
    parser.add_argument(
        "--chunk-chars",
        type=int,
        default=3000,
        help="Approximate max chars per text chunk (default: 3000)",
    )
    parser.add_argument(
        "--overlap-chars",
        type=int,
        default=250,
        help="Approximate chunk overlap in chars (default: 250)",
    )
    args = parser.parse_args()

    input_path = Path(args.input).expanduser().resolve()
    outdir = Path(args.outdir).expanduser().resolve()
    outdir.mkdir(parents=True, exist_ok=True)

    if not input_path.exists():
        print(f"error: input does not exist: {input_path}", file=sys.stderr)
        return 2
    if input_path.suffix.lower() != ".docx":
        print(f"error: input is not a .docx file: {input_path}", file=sys.stderr)
        return 2

    try:
        from docx import Document  # type: ignore
    except Exception as exc:  # pragma: no cover
        print(f"error: python-docx import failed: {exc}", file=sys.stderr)
        return 2

    document = Document(str(input_path))
    segments = list(iter_segments(document))
    chunks = list(iter_chunks(segments, args.chunk_chars, args.overlap_chars))

    with (outdir / "segments.jsonl").open("w", encoding="utf-8") as f:
        for row in segments:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")

    with (outdir / "chunks.jsonl").open("w", encoding="utf-8") as f:
        for row in chunks:
            f.write(json.dumps(row, ensure_ascii=False) + "\n")

    paragraph_count = sum(1 for row in segments if row["kind"] == "paragraph")
    table_row_count = sum(1 for row in segments if row["kind"] == "table_row")
    total_chars = sum(row["char_count"] for row in segments)
    manifest = {
        "input_path": str(input_path),
        "output_dir": str(outdir),
        "generated_at_utc": dt.datetime.now(dt.timezone.utc).isoformat(),
        "file_size_bytes": os.path.getsize(input_path),
        "segment_count": len(segments),
        "paragraph_count": paragraph_count,
        "table_row_count": table_row_count,
        "total_text_chars": total_chars,
        "chunk_count": len(chunks),
        "chunk_chars": args.chunk_chars,
        "overlap_chars": args.overlap_chars,
    }

    with (outdir / "manifest.json").open("w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2, ensure_ascii=False)

    print(json.dumps(manifest, ensure_ascii=False, separators=(",", ":")))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
