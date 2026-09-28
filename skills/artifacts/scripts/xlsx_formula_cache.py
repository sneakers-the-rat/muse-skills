"""Copy calculated values into an OOXML package without round-tripping its model."""

import io
import posixpath
import zipfile
from xml.dom import minidom
from xml.etree import ElementTree as ET

from openpyxl.utils.cell import (
    coordinate_to_tuple,
    get_column_letter,
    range_boundaries,
    range_to_tuple,
)

SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def sheet_names(archive):
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    return {
        sheet.attrib["name"]
        for sheet in workbook.findall(f"{{{SHEET_NS}}}sheets/{{{SHEET_NS}}}sheet")
    }


def worksheet_parts(archive):
    relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    targets = {
        rel.attrib["Id"]: posixpath.normpath(
            posixpath.join("xl", rel.attrib["Target"])
        ).lstrip("/")
        for rel in relationships
        if rel.attrib.get("Type", "").endswith("/worksheet")
    }
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    return {
        sheet.attrib["name"]: targets[sheet.attrib[f"{{{REL_NS}}}id"]]
        for sheet in workbook.findall(f"{{{SHEET_NS}}}sheets/{{{SHEET_NS}}}sheet")
        if sheet.attrib[f"{{{REL_NS}}}id"] in targets
    }


def child(element, local_name, namespace=SHEET_NS):
    return next(
        (
            node
            for node in element.childNodes
            if node.nodeType == node.ELEMENT_NODE
            and node.namespaceURI == namespace
            and node.localName == local_name
        ),
        None,
    )


def calculated_cells(archive, part, shared_strings):
    result = {}
    formulas = {}
    for cell in ET.fromstring(archive.read(part)).iter(f"{{{SHEET_NS}}}c"):
        formula = cell.find(f"{{{SHEET_NS}}}f")
        if formula is not None:
            formulas[cell.attrib["r"]] = formula.attrib
        value = cell.find(f"{{{SHEET_NS}}}v")
        kind = cell.get("t", "n")
        text = value.text if value is not None else None
        if kind == "s":
            text = shared_strings[int(text)]
            kind = "str"
        elif kind == "inlineStr":
            inline = cell.find(f"{{{SHEET_NS}}}is")
            text = "".join(inline.itertext()) if inline is not None else ""
            kind = "str"
        elif kind not in {"n", "str", "b", "e"}:
            raise ValueError(f"unsupported calculated cell type {kind!r}")
        if text is None and kind != "str" and formula is not None:
            raise ValueError(f"no calculated value for {part}:{cell.get('r')}")
        result[cell.attrib["r"]] = kind, text or ""
    return result, formulas


def chart_caches(archive):
    caches = {}
    for name in chart_parts(archive):
        root = ET.fromstring(archive.read(name))
        if root.tag != f"{{{CHART_NS}}}chartSpace":
            continue
        for parent in root.iter():
            for kind in ("num", "str", "multiLvlStr"):
                reference = parent.find(f"{{{CHART_NS}}}{kind}Ref")
                if reference is None:
                    continue
                formula = reference.find(f"{{{CHART_NS}}}f")
                cache = reference.find(f"{{{CHART_NS}}}{kind}Cache")
                if formula is not None and cache is not None:
                    key = parent.tag, range_to_tuple(formula.text)
                    caches[key] = kind, ET.tostring(cache)
    return caches


def chart_parts(archive):
    content_types = ET.fromstring(archive.read("[Content_Types].xml"))
    return {
        part.attrib["PartName"].lstrip("/")
        for part in content_types
        if part.get("ContentType")
        == "application/vnd.openxmlformats-officedocument.drawingml.chart+xml"
    }


def refresh_chart_caches(data, caches):
    document = minidom.parseString(data)
    if document.documentElement.namespaceURI != CHART_NS:
        document.unlink()
        return data
    changed = False
    for kind in ("num", "str", "multiLvlStr"):
        for reference in document.getElementsByTagNameNS(CHART_NS, f"{kind}Ref"):
            formula = child(reference, "f", CHART_NS)
            if formula is None or formula.firstChild is None:
                raise ValueError("chart reference has no source formula")
            text = formula.firstChild.data
            key = (
                f"{{{CHART_NS}}}{reference.parentNode.localName}",
                range_to_tuple(text),
            )
            if key not in caches:
                raise ValueError(f"recalculation did not refresh chart source {text!r}")
            calculated_kind, cache_bytes = caches[key]
            if kind != calculated_kind and "multiLvlStr" in {kind, calculated_kind}:
                raise ValueError(
                    f"recalculation changed chart category levels for {text!r}"
                )
            cache_document = minidom.parseString(cache_bytes)
            cache = document.importNode(cache_document.documentElement, True)
            cache_document.unlink()
            old = child(reference, f"{kind}Cache", CHART_NS)
            if old is not None:
                # Display formatting belongs to the original chart, not Calc.
                old_format = child(old, "formatCode", CHART_NS)
                new_format = child(cache, "formatCode", CHART_NS)
                if old_format is not None and calculated_kind == "num":
                    kept = document.importNode(old_format, True)
                    if new_format is not None:
                        cache.replaceChild(kept, new_format)
                    else:
                        cache.insertBefore(kept, cache.firstChild)
                reference.replaceChild(cache, old)
            else:
                reference.insertBefore(cache, formula.nextSibling)
            # Excel cache carriers distinguish text from numbers. Calc may
            # resolve an uncached category range as text, even when the
            # original writer emitted numRef; keep the range and use its
            # calculated cache type without changing the surrounding chart.
            prefix = f"{reference.prefix}:" if reference.prefix else ""
            document.renameNode(reference, CHART_NS, f"{prefix}{calculated_kind}Ref")
            changed = True
    result = document.toxml(encoding="UTF-8") if changed else data
    document.unlink()
    return result


def formula_coordinates(document, calculated_formulas):
    """Include every cached result in a fixed array, not just its anchor."""
    coordinates = set()
    for cell in document.getElementsByTagNameNS(SHEET_NS, "c"):
        formula = child(cell, "f")
        if formula is None:
            continue
        coordinate = cell.getAttribute("r")
        if coordinate not in calculated_formulas:
            raise ValueError(f"recalculation removed formula {coordinate}")
        kind = formula.getAttribute("t")
        if kind in {"", "normal", "shared"}:
            coordinates.add(coordinate)
            continue
        if kind != "array":
            raise ValueError(f"unsupported formula type {kind!r} at {coordinate}")
        reference = formula.getAttribute("ref")
        calculated = calculated_formulas[coordinate]
        if calculated.get("t") != "array" or range_boundaries(
            calculated["ref"]
        ) != range_boundaries(reference):
            raise ValueError(f"recalculation changed the array range at {coordinate}")
        left, top, right, bottom = range_boundaries(reference)
        coordinates.update(
            f"{get_column_letter(col)}{row}"
            for row in range(top, bottom + 1)
            for col in range(left, right + 1)
        )
    return coordinates


def insert_array_cells(document, additions):
    """Materialize cached array results omitted by the original writer."""
    sheet_data = child(document.documentElement, "sheetData")
    rows = list(sheet_data.getElementsByTagNameNS(SHEET_NS, "row"))
    prefix = f"{sheet_data.prefix}:" if sheet_data.prefix else ""
    row_index = 0
    for row_number, cells in sorted(additions.items()):
        while (
            row_index < len(rows)
            and int(rows[row_index].getAttribute("r")) < row_number
        ):
            row_index += 1
        following = rows[row_index] if row_index < len(rows) else None
        if following is not None and int(following.getAttribute("r")) == row_number:
            row = following
        else:
            row = document.createElementNS(SHEET_NS, f"{prefix}row")
            row.setAttribute("r", str(row_number))
            sheet_data.insertBefore(row, following)
        original_cells = list(row.getElementsByTagNameNS(SHEET_NS, "c"))
        cell_index = 0
        for column, cell in sorted(cells):
            while (
                cell_index < len(original_cells)
                and coordinate_to_tuple(original_cells[cell_index].getAttribute("r"))[1]
                < column
            ):
                cell_index += 1
            following_cell = (
                original_cells[cell_index]
                if cell_index < len(original_cells)
                else child(row, "extLst")
            )
            row.insertBefore(cell, following_cell)
        # This optional optimization hint no longer describes the row's cells.
        if row.hasAttribute("spans"):
            row.removeAttribute("spans")
    dimension = child(document.documentElement, "dimension")
    if dimension is not None:
        bounds = range_boundaries(dimension.getAttribute("ref"))
        left, top, right, bottom = bounds
        for row_number, cells in additions.items():
            top, bottom = min(top, row_number), max(bottom, row_number)
            left = min(left, min(column for column, _ in cells))
            right = max(right, max(column for column, _ in cells))
        if (left, top, right, bottom) != bounds:
            dimension.setAttribute(
                "ref",
                f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}",
            )


def merge_formula_caches(original_bytes, calculated_path):
    """Keep original package parts; replace only formula cell type/value caches.

    Namespace declarations must survive even when they occur only in attribute
    values such as mc:Ignorable. minidom retains them, unlike ElementTree's
    namespace-rewriting serializer. The calculated package is never delivered.
    """
    output = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(original_bytes)) as original,
        zipfile.ZipFile(calculated_path) as calculated,
        zipfile.ZipFile(output, "w") as merged,
    ):
        source_parts = worksheet_parts(original)
        result_parts = worksheet_parts(calculated)
        if sheet_names(original) != sheet_names(calculated):
            raise ValueError("recalculation changed the sheet inventory")
        shared_strings = []
        if "xl/sharedStrings.xml" in calculated.namelist():
            shared_strings = [
                "".join(
                    text.text or ""
                    for text in node.findall(f"{{{SHEET_NS}}}t")
                    + node.findall(f"{{{SHEET_NS}}}r/{{{SHEET_NS}}}t")
                )
                for node in ET.fromstring(calculated.read("xl/sharedStrings.xml"))
            ]
        replacements = {}
        for sheet_name, source_part in source_parts.items():
            values, formulas = calculated_cells(
                calculated, result_parts[sheet_name], shared_strings
            )
            document = minidom.parseString(original.read(source_part))
            changed = False
            coordinates = formula_coordinates(document, formulas)
            cells = {
                cell.getAttribute("r"): cell
                for cell in document.getElementsByTagNameNS(SHEET_NS, "c")
            }
            additions = {}
            for coordinate in sorted(coordinates, key=coordinate_to_tuple):
                if coordinate not in values:
                    raise ValueError(
                        f"recalculation removed formula {sheet_name}!{coordinate}"
                    )
                cell = cells.get(coordinate)
                if cell is None:
                    prefix = (
                        f"{document.documentElement.prefix}:"
                        if document.documentElement.prefix
                        else ""
                    )
                    cell = document.createElementNS(SHEET_NS, f"{prefix}c")
                    cell.setAttribute("r", coordinate)
                    row, column = coordinate_to_tuple(coordinate)
                    additions.setdefault(row, []).append((column, cell))
                formula = child(cell, "f")
                kind, text = values[coordinate]
                if kind != "str" and not text:
                    raise ValueError(
                        f"no calculated value for {sheet_name}!{coordinate}"
                    )
                cell.setAttribute("t", kind)
                for name in ("v", "is"):
                    old = child(cell, name)
                    if old is not None:
                        cell.removeChild(old)
                prefix = f"{cell.prefix}:" if cell.prefix else ""
                value = document.createElementNS(SHEET_NS, f"{prefix}v")
                value.appendChild(document.createTextNode(text))
                cell.insertBefore(
                    value, formula.nextSibling if formula else cell.firstChild
                )
                changed = True
            if additions:
                insert_array_cells(document, additions)
            if changed:
                replacements[source_part] = document.toxml(encoding="UTF-8")
            document.unlink()
        caches = chart_caches(calculated)
        charts = chart_parts(original)
        for member in original.infolist():
            data = replacements.get(member.filename, original.read(member))
            if member.filename in charts:
                data = refresh_chart_caches(data, caches)
            merged.writestr(member, data)
        merged.comment = original.comment
    return output.getvalue()
