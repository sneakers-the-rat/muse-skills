#!/usr/bin/env python3
"""Build an image-based PPTX deck from validated slide PNGs.

Usage: build_pptx.py <artifact-slug> [title] [--project-dir <dir>]

Reads the rendered, validated page PNGs from
  <project-dir>/.src/validate/page-*.png
and writes a full-bleed, one-image-per-slide deck to
  <project-dir>/<slug>.pptx

`--project-dir` is the `project_dir` the build task names. Pass it whenever the
build task gives one: an artifact built under a goal lives in that goal's
`files/` directory, not in `your_files`, and a hardcoded path would write the
deck somewhere nothing reads. It defaults to
`$JARVIS_HOME/workspace/your_files/<slug>` for a build task that names none.

Each slide is sized to the PNG aspect ratio so the image fills the slide without
distortion. Requires python-pptx
(`python3 -m pip install --break-system-packages python-pptx`).
"""

import json
import os
import re
import struct
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu


def slide_titles(project_dir, count):
    """Per-slide titles from the authored deck manifest, deck order.

    The PNG slides carry no text layer, so the titles become speaker notes:
    the one place a screen reader or presenter view can still read the deck.
    Fail-open: a missing or unparseable manifest (plain PDF-only builds have
    none) yields no notes, never a failed export.
    """
    manifest = project_dir / ".src" / "slides" / "deck.json"
    try:
        slides = json.loads(manifest.read_text()).get("slides", [])
    except (OSError, ValueError):
        return [None] * count
    titles = [
        entry.get("page_title") if isinstance(entry, dict) else None
        for entry in slides
    ]
    titles += [None] * max(0, count - len(titles))
    return titles[:count]


def page_number(path):
    match = re.search(r"page-(\d+)\.png$", path.name)
    if not match:
        raise SystemExit(f"unexpected PNG filename {path.name}; expected page-N.png")
    return int(match.group(1))


def png_size(path):
    # Read width/height from the PNG IHDR chunk (stdlib only, no Pillow).
    with open(path, "rb") as fh:
        header = fh.read(24)
    if header[:8] != b"\x89PNG\r\n\x1a\n":
        raise SystemExit(f"{path.name} is not a PNG")
    return struct.unpack(">II", header[16:24])


def split_project_dir(argv):
    """Pull `--project-dir <dir>` out of argv, returning (rest, dir_or_None).

    Only the LAST occurrence wins, and a value is never taken from a token that
    is itself an option. A deck title is a positional here, so a title that
    happens to read `--project-dir` would otherwise swallow the real flag as
    its value and build into a directory named after the option.
    """
    rest = []
    project_dir = None
    index = 0
    while index < len(argv):
        arg = argv[index]
        if arg == "--project-dir" and index + 1 < len(argv) and argv[index + 1].startswith("--"):
            # The next token is an option, so this one is a positional title.
            rest.append(arg)
            index += 1
            continue
        if arg == "--project-dir":
            if index + 1 >= len(argv):
                raise SystemExit("--project-dir needs a directory")
            project_dir = argv[index + 1]
            index += 2
            continue
        if arg.startswith("--project-dir="):
            project_dir = arg.split("=", 1)[1]
            index += 1
            continue
        rest.append(arg)
        index += 1
    return rest, project_dir


def main(argv):
    argv, project_dir_arg = split_project_dir(argv)
    if len(argv) < 2:
        raise SystemExit("usage: build_pptx.py <slug> [title] [--project-dir <dir>]")
    jarvis_home = os.environ.get("JARVIS_HOME")
    if not jarvis_home:
        raise SystemExit("JARVIS_HOME is not set")

    slug = argv[1]
    title = argv[2] if len(argv) > 2 else slug.replace("-", " ").title()
    # The build task's `project_dir` wins. It is the only thing that knows a
    # goal document belongs under that goal rather than in the shared area.
    if project_dir_arg is not None:
        if not project_dir_arg.strip():
            raise SystemExit("--project-dir was given an empty directory")
        # The contract spells `project_dir` with a leading `~`, and the shell
        # leaves that literal inside quotes, so the value can arrive either
        # tilde-prefixed or still carrying `$JARVIS_HOME`. Expand both rather
        # than writing a directory named `~` or `$JARVIS_HOME`.
        project_dir = Path(os.path.expanduser(os.path.expandvars(project_dir_arg)))
    else:
        project_dir = Path(jarvis_home) / "workspace" / "your_files" / slug
    validate_dir = project_dir / ".src" / "validate"
    output_pptx = project_dir / f"{slug}.pptx"

    pngs = sorted(validate_dir.glob("page-*.png"), key=page_number)
    if not pngs:
        raise SystemExit(
            f"no validation PNGs found in {validate_dir}; run the slide/PDF validation render first"
        )

    # Size the slide to the PNG aspect ratio. Decks default to 16:9, but the user
    # may request another ratio; the PNGs are the source of truth. Width is fixed
    # at 13.333in and height is derived so images are full-bleed without distortion.
    px_w, px_h = png_size(pngs[0])
    slide_w = Emu(12192000)  # 13.333in
    slide_h = Emu(round(12192000 * px_h / px_w))

    prs = Presentation()
    prs.core_properties.title = title
    prs.slide_width = slide_w
    prs.slide_height = slide_h
    # Find the Blank layout by name; fail loudly rather than silently picking a
    # layout that carries placeholder shapes.
    blank = next(
        (layout for layout in prs.slide_layouts if layout.name == "Blank"), None
    )
    if blank is None:
        raise SystemExit("no 'Blank' slide layout found in the default template")
    titles = slide_titles(project_dir, len(pngs))
    for png, note in zip(pngs, titles):
        slide = prs.slides.add_slide(blank)
        # Slide is sized to the PNG aspect ratio, so width+height fill exactly
        # with no stretch. python-pptx older than 0.6.18 rejects Path; str() is
        # portable.
        slide.shapes.add_picture(str(png), 0, 0, width=slide_w, height=slide_h)
        if note:
            slide.notes_slide.notes_text_frame.text = note
    prs.save(str(output_pptx))
    print(f"wrote {output_pptx} ({len(pngs)} slides)")


if __name__ == "__main__":
    main(sys.argv)
